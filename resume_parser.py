import re
from pathlib import Path
from typing import Dict, Any, List
import pypdf
import docx

# Comprehensive Tech & Professional Skill Dictionary
COMMON_SKILLS = [
    # Languages & Scripting
    "Python", "JavaScript", "TypeScript", "Java", "C++", "C#", "Go", "Rust", "PHP", 
    "Ruby", "Swift", "Kotlin", "Scala", "R", "SQL", "HTML", "CSS", "Bash", "Shell",
    "PowerShell", "Shell Scripting",

    # IT Support, Service Desk & Enterprise Administration
    "ServiceNow", "Salesforce", "Active Directory", "Azure Active Directory", "Microsoft Entra ID",
    "Entra ID", "Intune", "Microsoft Intune", "Incident Management", "Problem Management",
    "Change Management", "SLA Compliance", "Service Desk", "Technical Support", "Helpdesk",
    "ITIL", "ITSM", "Virtual Desktop Infrastructure", "VDI", "Citrix", "VMware", "Hyper-V",
    "Virtualization", "Performance Monitoring", "Exchange Admin Portal", "Microsoft 365",
    "Office 365", "M365", "Windows", "Windows Server", "Linux", "Unix", "Ubuntu", "CentOS",
    "RHEL", "Red Hat", "Dell PowerEdge", "Root Cause Analysis", "RCA", "Microsoft Dynamics",
    "Dynamics 365", "Admin Portal", "Remote Desktop", "RDP", "AnyDesk", "TeamViewer",
    "Network Troubleshooting", "Hardware Troubleshooting", "LAN/WAN", "DNS", "DHCP",
    "TCP/IP", "VPN", "Firewall", "Zendesk", "Freshdesk", "Confluence", "Jira",
    
    # Frontend & Fullstack
    "React", "React Native", "Vue.js", "Vue", "Angular", "Next.js", "Nuxt.js", 
    "Svelte", "Redux", "Tailwind CSS", "Bootstrap", "Webpack", "Vite", "GraphQL", 
    "REST API", "Node.js", "Express.js", "Django", "Flask", "FastAPI", "Spring Boot",
    "Ruby on Rails", "ASP.NET", "Laravel",
    
    # Data & AI / ML
    "Machine Learning", "Deep Learning", "Artificial Intelligence", "LLM", "NLP",
    "Computer Vision", "TensorFlow", "PyTorch", "Keras", "Scikit-Learn", "Pandas",
    "NumPy", "OpenCV", "Hugging Face", "LangChain", "LlamaIndex", "Generative AI",
    "Prompt Engineering", "Data Engineering", "Apache Spark", "Airflow", "Kafka",
    "Snowflake", "Databricks", "BigQuery", "Tableau", "Power BI",
    
    # Cloud & DevOps
    "AWS", "Amazon Web Services", "Azure", "Google Cloud", "GCP", "Docker", 
    "Kubernetes", "Terraform", "CI/CD", "GitHub Actions", "GitLab CI", "Jenkins", 
    "Ansible", "Prometheus", "Grafana", "Nginx", "Microservices",
    
    # Databases
    "PostgreSQL", "MySQL", "MongoDB", "Redis", "Elasticsearch", "DynamoDB", 
    "SQLite", "Cassandra", "Supabase", "Firebase",
    
    # Tools & Practices
    "Git", "GitHub", "GitLab", "Agile", "Scrum", "TDD", "Unit Testing",
    "System Design", "OOP", "Software Architecture", "Cybersecurity", "Penetration Testing"
]

def extract_text_from_pdf(file_bytes_or_path) -> str:
    reader = pypdf.PdfReader(file_bytes_or_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def extract_text_from_docx(file_bytes_or_path) -> str:
    doc = docx.Document(file_bytes_or_path)
    full_text = []
    try:
        from docx.oxml.text.paragraph import CT_P
        from docx.oxml.table import CT_Tbl
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        for block in doc.element.body:
            if isinstance(block, CT_P):
                p = Paragraph(block, doc)
                if p.text:
                    full_text.append(p.text)
            elif isinstance(block, CT_Tbl):
                t = Table(block, doc)
                for row in t.rows:
                    seen_cells = set()
                    row_cells = []
                    for cell in row.cells:
                        if cell._tc not in seen_cells:
                            seen_cells.add(cell._tc)
                            ctext = cell.text.strip().replace('\r', ' ')
                            if ctext:
                                row_cells.append(ctext)
                    if row_cells:
                        full_text.append(" | ".join(row_cells))
    except Exception:
        # Fallback to paragraph and table iteration
        full_text = [p.text for p in doc.paragraphs if p.text]
        for t in doc.tables:
            for row in t.rows:
                row_cells = [c.text.strip().replace('\n', ' ') for c in row.cells if c.text.strip()]
                if row_cells:
                    full_text.append(" | ".join(row_cells))

    return "\n".join(full_text)

def extract_email(text: str) -> str:
    match = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', text)
    return match.group(0) if match else ""

def extract_phone(text: str) -> str:
    match = re.search(r'(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}', text)
    return match.group(0).strip() if match else ""

def extract_location(text: str) -> str:
    lines = [l.strip() for l in text.split('\n') if l.strip()][:12]
    countries = r'India|USA|United States|UK|United Kingdom|Canada|Germany|Australia|Singapore|France|Netherlands|Ireland|UAE|Dubai|Japan|Poland|Spain|Remote'
    cities = r'Pune|Mumbai|Bangalore|Bengaluru|Delhi|New Delhi|Hyderabad|Chennai|Kolkata|Noida|Gurgaon|Gurugram|Ahmedabad|Aundh|San Francisco|New York|Seattle|Austin|London|Berlin|Toronto|Vancouver'
    
    for line in lines:
        parts = re.split(r'\||•|·|\t', line)
        for part in parts:
            p = part.strip()
            if '@' in p or re.search(r'\b\d{7,}\b', p):
                continue
            # Pattern 1: City, Country / Pincode (e.g. Aundh, India 411007)
            m = re.search(rf'([A-Za-z\s]+,\s*(?:{countries})\b(?:\s*\d{{5,6}})?)', p, re.I)
            if m:
                return m.group(1).strip()
            # Pattern 2: City, US State Zip (e.g. San Francisco, CA 94105)
            m2 = re.search(r'([A-Za-z\s]+,\s*[A-Z]{2}\b(?:\s*\d{5})?)', p)
            if m2 and len(m2.group(1).split()) <= 6:
                return m2.group(1).strip()
            # Pattern 3: City, State/Region (e.g. Pune, Maharashtra)
            m3 = re.search(rf'(\b(?:{cities})\b(?:,\s*[A-Za-z\s]+)?)', p, re.I)
            if m3 and len(m3.group(1).split()) <= 4:
                return m3.group(1).strip()
            # Pattern 4: Standalone country/Remote
            m4 = re.search(rf'^(?:{countries})$', p, re.I)
            if m4:
                return m4.group(0).strip()
    return ""

def extract_links(text: str) -> Dict[str, str]:
    links = {}
    linkedin = re.search(r'(https?://(?:www\.)?linkedin\.com/in/[\w-]+)', text, re.I)
    if linkedin:
        links['linkedin'] = linkedin.group(0)
    github = re.search(r'(https?://(?:www\.)?github\.com/[\w-]+)', text, re.I)
    if github:
        links['github'] = github.group(0)
    portfolio = re.search(r'(https?://(?:www\.)?[\w-]+\.(?:com|dev|io|me|org)[\w/-]*)', text, re.I)
    if portfolio and 'linkedin' not in portfolio.group(0) and 'github' not in portfolio.group(0):
        links['portfolio'] = portfolio.group(0)
    return links

def extract_skills_from_section(skills_raw: str) -> List[str]:
    raw_skills = []
    lines = skills_raw.split('\n')
    for line in lines:
        tokens = re.split(r'\||•|·|\u2022|\uf0b7|\ufffd|\t|;', line)
        for t in tokens:
            cleaned = re.sub(r'^[\s\u2022\uf0b7\ufffd\*\-\•\·\>\#\:\.\,\;\|\d\)]+', '', t)
            cleaned = re.sub(r'[\s\.\,\;\|]+$', '', cleaned).strip()
            if not cleaned or len(cleaned) < 2 or len(cleaned) > 50 or len(cleaned.split()) > 5:
                continue
            # Filter noise words
            if cleaned.lower() in ["skills", "technical skills", "competencies", "tools", "technologies"]:
                continue
            raw_skills.append(cleaned)
    return raw_skills

def extract_skills(text: str) -> List[str]:
    found_skills = set()
    lower_map = {}
    
    # 1. First extract skills listed under explicit SKILLS sections (including tables)
    sections = extract_sections(text)
    if sections.get("skills_raw"):
        explicit_items = extract_skills_from_section(sections["skills_raw"])
        for item in explicit_items:
            found_skills.add(item)
            lower_map[item.lower()] = item
            
    # 2. Check full text against COMMON_SKILLS dictionary
    for skill in COMMON_SKILLS:
        pattern = rf'\b{re.escape(skill)}\b'
        if re.search(pattern, text, re.IGNORECASE):
            if skill.lower() not in lower_map:
                found_skills.add(skill)
                lower_map[skill.lower()] = skill
            
    # Sort alphabetically by case-insensitive key
    return sorted(list(found_skills), key=lambda x: x.lower())

def extract_sections(text: str) -> Dict[str, str]:
    sections = {
        "summary": "",
        "experience": "",
        "education": "",
        "projects": "",
        "skills_raw": ""
    }
    
    lines = text.split("\n")
    current_sec = None
    sec_lines = {
        "summary": [],
        "experience": [],
        "education": [],
        "projects": [],
        "skills_raw": []
    }
    
    for line in lines:
        clean = line.strip().lower()
        if not clean:
            continue
            
        if re.match(r'^(summary|professional summary|executive summary|about me|profile|overview)\b', clean):
            current_sec = "summary"
            continue
        elif re.match(r'^(experience|work experience|employment history|work history|professional experience|career history)\b', clean):
            current_sec = "experience"
            continue
        elif re.match(r'^(education|academic background|academics|qualifications|educational background)\b', clean):
            current_sec = "education"
            continue
        elif re.match(r'^(projects|personal projects|technical projects|key projects|academic projects)\b', clean):
            current_sec = "projects"
            continue
        elif re.match(r'^(skills|technical skills|core competencies|key skills|technologies|tools & technologies|tools|competencies|areas of expertise)\b', clean):
            current_sec = "skills_raw"
            continue
        elif re.match(r'^(languages|certifications|certificates|awards|honors|publications|hobbies|interests|hobbies & interests|activities|extracurricular|volunteer|references)\b', clean):
            current_sec = None
            continue
            
        if current_sec and current_sec in sec_lines:
            sec_lines[current_sec].append(line)
            
    sections["summary"] = "\n".join(sec_lines["summary"][:20]).strip()
    sections["experience"] = "\n".join(sec_lines["experience"][:60]).strip()
    sections["education"] = "\n".join(sec_lines["education"][:25]).strip()
    sections["projects"] = "\n".join(sec_lines["projects"][:40]).strip()
    sections["skills_raw"] = "\n".join(sec_lines["skills_raw"][:40]).strip()
    return sections

def parse_resume_text(raw_text: str, filename: str = "Pasted Resume") -> Dict[str, Any]:
    """Parse resume from raw string text directly."""
    lines = [l.strip() for l in raw_text.split('\n') if l.strip()]
    guessed_name = "Candidate"
    
    for l in lines[:6]:
        lower_l = l.lower()
        if any(h in lower_l for h in ["resume", "curriculum vitae", "cv", "portfolio", "profile", "contact", "summary"]):
            continue
        if "@" in l or re.search(r'\b\d{6,}\b', l):
            continue
        if len(l.split()) <= 5 and not any(ch in l for ch in ["|", ":", "/", "\\", "@"]):
            guessed_name = l.title() if l.isupper() else l
            break
            
    email = extract_email(raw_text)
    phone = extract_phone(raw_text)
    location = extract_location(raw_text)
    links = extract_links(raw_text)
    sections = extract_sections(raw_text)
    skills = extract_skills(raw_text)
    
    return {
        "name": guessed_name[:60],
        "email": email,
        "phone": phone,
        "location": location,
        "links": links,
        "skills": skills,
        "summary": sections.get("summary", ""),
        "experience": sections.get("experience", ""),
        "education": sections.get("education", ""),
        "projects": sections.get("projects", ""),
        "raw_text": raw_text
    }

def parse_resume(file_obj_or_path, filename: str) -> Dict[str, Any]:
    raw_text = ""
    lower_fn = filename.lower()
    
    if lower_fn.endswith('.pdf'):
        raw_text = extract_text_from_pdf(file_obj_or_path)
    elif lower_fn.endswith('.docx'):
        raw_text = extract_text_from_docx(file_obj_or_path)
    elif lower_fn.endswith('.txt'):
        if hasattr(file_obj_or_path, 'read'):
            raw = file_obj_or_path.read()
            raw_text = raw.decode('utf-8', errors='ignore') if isinstance(raw, bytes) else str(raw)
        else:
            with open(file_obj_or_path, 'r', encoding='utf-8', errors='ignore') as f:
                raw_text = f.read()
    else:
        # Fallback treat as text
        if hasattr(file_obj_or_path, 'read'):
            raw = file_obj_or_path.read()
            raw_text = raw.decode('utf-8', errors='ignore') if isinstance(raw, bytes) else str(raw)
            
    return parse_resume_text(raw_text, filename)
